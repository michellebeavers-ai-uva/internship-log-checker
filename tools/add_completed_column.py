"""Add a 'Completed' column (Yes dropdown) to every course tab of the A&S Planning Log.

  python tools/add_completed_column.py "A_S_Planning_Log.xlsx" "A_S_Planning_Log_with_Completed.xlsx"

Always open the result in Excel and compare it with your original before sharing it with students.
The checker finds the column by its header text ("Completed"), wherever it sits.
"""
import sys
from copy import copy
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

HEADER = "Completed"


def add_column(src, dst):
    wb = load_workbook(src)
    changed = []
    for ws in wb.worksheets:
        if not ws.title.strip().isdigit():
            continue
        hr = next((r for r in range(1, 20) if str(ws.cell(r, 1).value or "").strip().lower() == "type"), None)
        if hr is None:
            continue
        headers = {str(c.value or "").strip().lower(): c.column for c in ws[hr]}
        if HEADER.lower() in headers:
            continue
        col = max(c.column for c in ws[hr] if c.value is not None) + 1
        letter = ws.cell(hr, col).column_letter
        left = ws.cell(hr, col - 1)
        ws.cell(hr, col).value = HEADER
        for attr in ("font", "fill", "border", "alignment"):
            setattr(ws.cell(hr, col), attr, copy(getattr(left, attr)))
        ws.column_dimensions[letter].width = 12
        dv = DataValidation(type="list", formula1='"Yes"', allow_blank=True, showDropDown=False)
        dv.promptTitle, dv.prompt = "Completed", "Choose Yes once you have finished this experience."
        dv.showInputMessage = True
        ws.add_data_validation(dv)
        for r in range(hr + 1, ws.max_row + 1):
            if str(ws.cell(r, 1).value or "").strip().upper() in ("EMBEDDED", "ALIGNED", "CHOICE"):
                dv.add(ws.cell(r, col))
                src_cell = ws.cell(r, col - 1)
                for attr in ("font", "border", "alignment"):
                    setattr(ws.cell(r, col), attr, copy(getattr(src_cell, attr)))
        changed.append(ws.title)
    wb.save(dst)
    return changed


if __name__ == "__main__":
    print("Added Completed column to:", ", ".join(add_column(sys.argv[1], sys.argv[2])))
