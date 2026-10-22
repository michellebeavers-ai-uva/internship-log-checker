"""Builds sample/sample_student_workbook.xlsx: the standard planning log (tools/make_standard_log.py), filled with INVENTED data."""
import os, sys, tempfile
from datetime import date
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from tools.make_standard_log import build as build_standard

TEMPLATE, OUT = sys.argv[1], sys.argv[2]
tmp = os.path.join(tempfile.mkdtemp(), "with_col.xlsx")
build_standard(TEMPLATE, tmp)
wb = load_workbook(tmp)
PLACEHOLDER = "Copy and paste choice hours here. No activity may be more than three hours"


def done_col(ws):
    return next(c.column for c in ws[5] if c.value == "Completed")


def fill(ws, row, d=None, loc=None, role=None, desc=None, hrs=None, comp=None, done=None):
    if d: ws.cell(row, 2).value = d
    if loc: ws.cell(row, 3).value = loc
    if role: ws.cell(row, 4).value = role
    if desc is not None: ws.cell(row, 5).value = desc
    if hrs is not None: ws.cell(row, 6).value = hrs
    if comp:
        for c in range(7, 13): ws.cell(row, c).value = None
        ws.cell(row, 6 + comp).value = "X"
    if done: ws.cell(row, done_col(ws)).value = "Yes"


for ws in wb.worksheets:
    if ws.title.strip().isdigit():
        ws["A2"].value = "Name: Pat Example      Placement(s): Invented Elementary"   # must never reach a report
        for r in range(6, ws.max_row + 1):
            if str(ws.cell(r, 1).value or "").upper() == "CHOICE":
                for c in range(2, 15): ws.cell(r, c).value = None

TEXT = {
    "7806": ["Facilitated a staff data meeting with the assistant principal and drafted a one-page summary of next steps",
             "Worked with the PTA president to plan a family night and wrote a reflection on community engagement barriers",
             "Observed a district budget work session and wrote notes connecting decisions to school priorities"],
    "7812": ["Led a grade-level team through an improvement-cycle review with the principal and recorded recommendations",
             "Interviewed two community partners about school outreach and summarized findings for the leadership team",
             "Shadowed the transportation director during bus route planning and reflected on operational tradeoffs"],
}
# Earlier courses (Summer 2026): finished, complete records
for code, loc in (("7806", "Agency"), ("7812", "Middle")):
    ws = wb[code]
    for r in (6, 7, 8):
        fill(ws, r, date(2026, 6, 10 + r), loc, "Participant", done=True)
    fill(ws, 10, date(2026, 7, 6), "High", "Leader", TEXT[code][0], 3, 5, done=True)
    fill(ws, 11, date(2026, 7, 9), "Elementary", "Participant", TEXT[code][1], 3, 4, done=True)
    fill(ws, 12, date(2026, 7, 14), "Central Office", "Observer", TEXT[code][2], 3, 3, done=True)

# Current term (Fall 2026): 7808 and 7813, partly done, with typical problems
ws = wb["7808"]
fill(ws, 6, date(2026, 9, 14), "Main Site", "Participant", done=True)
fill(ws, 7, date(2026, 9, 28), "Main Site", "Participant", done=True)
fill(ws, 8)                                                         # missing date/location/role
fill(ws, 10, date(2026, 10, 5), "Elementary", "Leader", "Help with budget", 2, 3, done=True)   # vague
fill(ws, 11, date(2026, 10, 8), "Central Office", "Observer",
     "Lead a culture change initiative for the whole school with the leadership team and create an action plan", 3, 5, done=True)
fill(ws, 12, date(2026, 10, 12), "Middle", "Participant", "Attended finance committee", 4, 6)  # over 3 hours, not done
fill(ws, 13, desc=PLACEHOLDER)                                      # untouched placeholder
ws = wb["7813"]
for r in (6, 7):
    fill(ws, r, date(2026, 10, 2 + r), "Main Site", "Participant", done=True)
fill(ws, 8, date(2026, 10, 10), "Main Site", "Participant")         # dated, but not marked complete
fill(ws, 10, date(2026, 11, 20), "High", "Participant",
     "Lead a culture change initiative for the whole school with the leadership team and create an action plan", 2, 5)
fill(ws, 11, desc=PLACEHOLDER)                                      # same placeholder in another tab
# 7813 row 12: copied word-for-word from a Choice Library sample (the checker should flag it)
from engine.parse import _read_library
lib = _read_library(wb)
sample = next(e for e in lib if e.course == "7813") if any(e.course == "7813" for e in lib) else lib[0]
fill(wb["7813"], 12, date(2026, 10, 25), "Middle", "Participant", sample.text, 1, sample.comp)
wb.save(OUT)
print("wrote", OUT)
