"""Builds sample/sample_2nd_med_workbook.xlsx: the 2nd MEd log (tools/make_second_med_log.py) filled with INVENTED data.
Scenario: Fall 2026, semester 3. Enrolled in 7805 and 7810, and completing the 7140 block of the '2nd Int Addl Reqs' tab."""
import os, re, sys, tempfile
from datetime import date
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from tools.make_second_med_log import build, NEW_TAB

TEMPLATE, OUT = sys.argv[1], sys.argv[2]
tmp = os.path.join(tempfile.mkdtemp(), "log.xlsx")
build(TEMPLATE, tmp)
wb = load_workbook(tmp)


def done_col(ws):
    return next(c.column for c in ws[5] if c.value == "Completed")


def fill(ws, row, d=None, loc=None, role=None, desc=None, hrs=None, comp=None, done=False):
    if d: ws.cell(row, 2).value = d
    if loc: ws.cell(row, 3).value = loc
    if role: ws.cell(row, 4).value = role
    if desc is not None: ws.cell(row, 5).value = desc
    if hrs is not None: ws.cell(row, 6).value = hrs
    if comp:
        for c in range(7, 13): ws.cell(row, c).value = None
        ws.cell(row, 6 + comp).value = "X"
    if done: ws.cell(row, done_col(ws)).value = "Yes"


def first_choice_row(ws):
    return next(r for r in range(6, ws.max_row + 1) if str(ws.cell(r, 1).value or "").upper() == "CHOICE")


def block_rows(ws, code):
    """Rows of one course's block on the combined tab: {'EMBEDDED': [...], 'ALIGNED': [...], 'CHOICE': [...]}."""
    out, current = {"EMBEDDED": [], "ALIGNED": [], "CHOICE": []}, None
    for r in range(6, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "").strip()
        m = re.match(r"^(\d{4})\s*[–-]", a)
        if m:
            current = m.group(1)
        elif current == code and a.upper() in out:
            out[a.upper()].append(r)
    return out


for ws in wb.worksheets:
    if ws.title.strip().isdigit() or ws.title == NEW_TAB:
        ws["A2"].value = "Name: Pat Example      Placement(s): Invented Elementary"   # must never reach a report
        for r in range(6, ws.max_row + 1):
            if str(ws.cell(r, 1).value or "").upper() == "CHOICE":
                for c in range(2, 16): ws.cell(r, c).value = None

# Finished earlier (semesters 1-2): dated before Fall 2026
for code in ("7806", "7812", "7807", "7811"):
    ws = wb[code]
    for r in range(6, first_choice_row(ws) - 1):
        if str(ws.cell(r, 1).value or "").upper() in ("EMBEDDED", "ALIGNED"):
            fill(ws, r, date(2026, 2, 10 + r), "Agency" if code == "7806" else "Middle", "Participant", done=True)

# Semester 3, Fall 2026
ws = wb["7805"]
for r, d in ((6, date(2026, 9, 15)), (7, date(2026, 9, 29)), (8, date(2026, 10, 6))):
    fill(ws, r, d, "Main Site", "Participant", done=True)
fill(ws, first_choice_row(ws), date(2026, 10, 9), "High", "Leader",
     "Led a hiring-committee debrief with two assistant principals and drafted recommendations for interview questions", 3, 3, done=True)
ws = wb["7810"]
for r, d in ((6, date(2026, 10, 14)), (7, date(2026, 11, 3)), (8, date(2026, 11, 10))):
    fill(ws, r, d, "Main Site", "Participant")

# The 7140 block of the combined tab, completed this semester alongside the enrolled courses
ws = wb[NEW_TAB]
b = block_rows(ws, "7140")
fill(ws, b["EMBEDDED"][0], date(2026, 10, 3), "Central Office", "Participant", done=True)      # embedded, 25h
fill(ws, b["ALIGNED"][0], date(2026, 10, 17), "Elementary", "Leader", done=True)               # aligned, 5h
fill(ws, b["CHOICE"][0], date(2026, 10, 18), "Middle", "Participant",
     "Analyzed attendance data with the counseling team and wrote a summary of operational recommendations", 2, 3, done=True)
ws["E7"].value = "Semester 3"                                                                   # the block's semester dropdown
wb.save(OUT)
print("wrote", OUT)
