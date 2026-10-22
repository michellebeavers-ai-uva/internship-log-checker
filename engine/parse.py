"""Read a student's planning-log workbook into plain Python objects."""
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

KINDS = ("EMBEDDED", "ALIGNED", "CHOICE")
COMP_LABELS = ["①", "②", "③", "④", "⑤", "⑥"]
DONE_VALUES = {"yes", "y", "x", "✓", "✔", "done", "true", "1", "complete", "completed"}


@dataclass
class Activity:
    course: str
    row: int
    kind: str
    date: date = None
    date_unreadable: bool = False
    location: str = ""
    role: str = ""
    description: str = ""
    hours: float = 0.0
    hours_unreadable: bool = False
    comps: list = field(default_factory=list)
    evidence: str = ""
    completed: bool = False      # only meaningful when the workbook has a "Completed" column
    sheet: str = ""              # tab the row lives on (the 2nd-MEd tab holds two courses' blocks)


@dataclass
class CourseTab:
    code: str
    activities: list = field(default_factory=list)
    sheet_choice_needed: float = None   # the number printed in the tab's CHOICE header
    has_completed_column: bool = False


@dataclass
class LibraryEntry:
    text: str
    comp: int = None
    course: str = None


@dataclass
class WorkbookModel:
    tabs: dict            # code -> CourseTab
    missing_tabs: list
    library: list
    has_completed_column: bool = False
    pathway_hint: str = None          # "second_med" when the Home tab says "Pathway: 2nd MEd"


def _clean(v):
    return "" if v is None else str(v).strip()


def parse_date(v):
    """-> (date or None, unreadable?)"""
    if v is None or _clean(v) == "":
        return None, False
    if isinstance(v, datetime):
        return v.date(), False
    if isinstance(v, date):
        return v, False
    s = _clean(v)
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%B %d, %Y", "%b %d, %Y", "%m-%d-%Y"):
        try:
            return datetime.strptime(s, fmt).date(), False
        except ValueError:
            pass
    return None, True


def parse_hours(v):
    """-> (hours, unreadable?)"""
    if v is None or _clean(v) == "":
        return 0.0, False
    if isinstance(v, (int, float)):
        return float(v), False
    m = re.search(r"\d+(?:\.\d+)?", _clean(v))
    if m:
        return float(m.group()), False
    return 0.0, True


def _col_map(ws, header_row):
    cols = {}
    for c in ws[header_row]:
        label = _clean(c.value).lower()
        if label:
            cols[label] = c.column
    return cols


def _find_header_row(ws):
    for r in range(1, 20):
        if _clean(ws.cell(r, 1).value).lower() == "type":
            return r
    return None


BLOCK_BAND = re.compile(r"^(\d{4})\s*[–-]")      # a band row such as "7140 – Leadership II (LEAD II)  |  32 hours"


def _read_rows(ws, default_code, tabs):
    """Read one sheet into `tabs` (code -> CourseTab).

    A normal course tab holds one course (`default_code`). The 2nd-MEd combined tab ("2nd Int Addl Reqs") passes
    default_code=None: its rows are credited to whichever course band ("7140 – ...", "7813 – ...") they sit under.
    """
    hr = _find_header_row(ws)
    if default_code is not None:
        tabs.setdefault(default_code, CourseTab(code=default_code))
    if hr is None:
        return
    cols = _col_map(ws, hr)
    c_date = cols.get("date", 2)
    c_loc = cols.get("location", 3)
    c_role = cols.get("role", 4)
    c_desc = cols.get("experience description", 5)
    c_hrs = cols.get("hrs", 6)
    c_comp = [cols.get(lbl, 7 + i) for i, lbl in enumerate(COMP_LABELS)]
    c_evid = cols.get("evidence / notes", 13)
    c_done = cols.get("completed")
    sheet = ws.title.strip()
    current = default_code

    for r in range(hr + 1, ws.max_row + 1):
        a = _clean(ws.cell(r, 1).value)
        if default_code is None:
            m = BLOCK_BAND.match(a)
            if m:
                current = m.group(1)
                tabs.setdefault(current, CourseTab(code=current))
                tabs[current].has_completed_column = c_done is not None
                continue
        if current is None:
            continue
        tab = tabs[current]
        tab.has_completed_column = c_done is not None
        if a.startswith("⭐") or "CHOICE HOURS" in a.upper():
            m = re.search(r"(\d+(?:\.\d+)?)\s*hours?\s*needed", a, re.I)
            if m:
                tab.sheet_choice_needed = float(m.group(1))
            continue
        kind = a.upper()
        if kind not in KINDS:
            continue
        act = Activity(course=current, row=r, kind=kind, sheet=sheet)
        act.date, act.date_unreadable = parse_date(ws.cell(r, c_date).value)
        act.location = _clean(ws.cell(r, c_loc).value)
        act.role = _clean(ws.cell(r, c_role).value)
        act.description = _clean(ws.cell(r, c_desc).value)
        act.hours, act.hours_unreadable = parse_hours(ws.cell(r, c_hrs).value)
        act.comps = [i + 1 for i, c in enumerate(c_comp)
                     if _clean(ws.cell(r, c).value).upper() == "X"]
        act.evidence = _clean(ws.cell(r, c_evid).value)
        if c_done:
            act.completed = _clean(ws.cell(r, c_done).value).lower() in DONE_VALUES
        blank = not (act.description or act.hours or act.date or act.location
                     or act.role or act.comps or act.date_unreadable or act.completed)
        if kind == "CHOICE" and blank:
            continue          # an unused choice row
        tab.activities.append(act)


def _read_library(wb):
    lib = []
    for ws in wb.worksheets:
        t = ws.title.lower()
        if "choice library-all" in t:
            for r in range(6, ws.max_row + 1):
                for i in range(6):
                    txt = _clean(ws.cell(r, 2 + i).value)
                    if len(txt) > 25:
                        lib.append(LibraryEntry(text=txt, comp=i + 1))
        elif "choice library-by" in t:
            course = None
            for r in range(4, ws.max_row + 1):
                a = _clean(ws.cell(r, 1).value)
                m = re.match(r"(\d{4})", a)
                if m:
                    course = m.group(1)
                focus = _clean(ws.cell(r, 2).value)
                txt = _clean(ws.cell(r, 3).value)
                if txt and focus[:1].isdigit():
                    lib.append(LibraryEntry(text=txt, comp=int(focus[0]), course=course))
    return lib


def _detect_pathway(wb):
    for ws in wb.worksheets:
        if "home" in ws.title.lower():
            for row in ws.iter_rows(min_row=1, max_row=6, max_col=3):
                for c in row:
                    if isinstance(c.value, str) and re.search(r"pathway:\s*2nd\s*med", c.value, re.I):
                        return "second_med"
    return None


def read_workbook(source, cfg):
    """source: a file path, bytes, or file-like object."""
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    wb = load_workbook(source, data_only=True)
    by_title = {ws.title.strip(): ws for ws in wb.worksheets}
    tabs, missing = {}, []
    for code in cfg["courses"]:
        ws = by_title.get(code)
        if ws is None:
            missing.append(code)
        else:
            _read_rows(ws, code, tabs)
    # 2nd-MEd log: the 7140 and 7813 experiences share one tab, in separate blocks
    for title, ws in by_title.items():
        if title.lower().startswith("2nd int"):
            found = {}
            _read_rows(ws, None, found)
            for code, tab in found.items():
                if code in cfg["courses"]:
                    tabs[code] = tab
                    if code in missing:
                        missing.remove(code)
    return WorkbookModel(tabs=tabs, missing_tabs=missing, library=_read_library(wb),
                         has_completed_column=any(t.has_completed_column for t in tabs.values()),
                         pathway_hint=_detect_pathway(wb))
